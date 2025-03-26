import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.chart import BarChart, Reference

df = pd.read_csv("test_data_new2.csv")

ok_count = df["Status"].value_counts().get("Ok")
failed_count = df["Status"].value_counts().get("Failed")

wb = load_workbook("report_data3.xlsx")
ws = wb.active

ws.append([])
ws.append(["Shrnutí výsledků"])
ws.append(["Počet OK testů", ok_count])
ws.append(["Počet Failed testů", failed_count])

chart = BarChart()
chart.title = "Výsledky testů"
chart.x_axis.title = "Status"
chart.y_axis.title = "Počet"

data = Reference(ws, min_col=2, min_row=10, max_row=11)
categories = Reference(ws, min_col=2, min_row=10, max_row=11)

chart.add_data(data, titles_from_data=False)
chart.set_categories(categories)

ws.add_chart(chart, "E17")

ws.column_dimensions["B"].width = 20
ws.column_dimensions["D"].width = 20
ws.column_dimensions["E"].width = 20

wb.save("report_final.xlsx")